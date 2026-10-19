"""Export the canonical evaluator's existing scores without recomputing metrics."""
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill


def export_evaluation_excel(aggregate, rows, path):
    """Validate every evaluation row and write the ranked scores and summary."""
    if not rows or len(rows) != aggregate['test_images']:
        raise ValueError('Per-image count must match the nonempty evaluated test set')
    scores = pd.DataFrame(rows)
    names = scores['image_name'] if 'image_name' in scores else scores['image']
    if names.isna().any() or not names.map(lambda name: isinstance(name, str) and bool(name.strip())).all():
        raise ValueError('Every test image must have a nonempty image name')
    if names.duplicated().any() or scores['image'].duplicated().any():
        raise ValueError('Duplicate per-image evaluation rows')
    for key in ('dice', 'iou'):
        values = scores[key].to_numpy(dtype=float)
        if not np.isfinite(values).all() or not ((values >= 0) & (values <= 1)).all():
            raise ValueError(f'{key} scores must be finite and between 0 and 1')
    for key in ('global_dice', 'global_iou', 'mean_per_image_dice', 'mean_per_image_iou'):
        value = aggregate[key]
        if not np.isfinite(value) or not 0 <= value <= 1:
            raise ValueError(f'Invalid aggregate metric: {key}')
    for column, key in (('dice', 'mean_per_image_dice'), ('iou', 'mean_per_image_iou')):
        if not np.isclose(scores[column].mean(), aggregate[key], rtol=0, atol=1e-12):
            raise ValueError(f'Per-image scores disagree with {key}')

    per_image = pd.DataFrame({
        'Image_Name': names, 'Dice_Score': scores['dice'],
        'Jaccard_Score': scores['iou'],
        'Pixel_Accuracy': scores['pixel_accuracy'], 'Loss': scores['loss'],
    }).sort_values('Dice_Score', ascending=False, kind='stable').reset_index(drop=True)
    per_image.insert(0, 'Rank', np.arange(1, len(per_image) + 1))
    summary = pd.DataFrame([
        ('Total Test Images', aggregate['test_images']),
        ('Global Dice', aggregate['global_dice']),
        ('Global Jaccard / Global IoU', aggregate['global_iou']),
        ('Mean Dice', aggregate['mean_per_image_dice']),
        ('Mean Jaccard', aggregate['mean_per_image_iou']),
        ('Median Dice', scores['dice'].median()),
        ('Median Jaccard', scores['iou'].median()),
        ('Best Dice', scores['dice'].max()),
        ('Worst Dice', scores['dice'].min()),
        ('Best Jaccard', scores['iou'].max()),
        ('Worst Jaccard', scores['iou'].min()),
    ], columns=['Metric', 'Value'])

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.stem + '.tmp.xlsx')
    try:
        with pd.ExcelWriter(temporary, engine='openpyxl') as writer:
            per_image.to_excel(writer, sheet_name='Per_Image_Metrics', index=False)
            summary.to_excel(writer, sheet_name='Summary', index=False)
            for sheet in writer.book.worksheets:
                sheet.sheet_view.showGridLines = False
                for row in sheet.iter_rows():
                    for cell in row:
                        cell.font = Font(name='Arial', size=10)
                        cell.alignment = Alignment(vertical='center', horizontal='right' if isinstance(cell.value, (int, float)) else 'left')
                        if isinstance(cell.value, (int, float)):
                            cell.number_format = '0.000000E+00' if 0 < abs(cell.value) < 1e-6 else '0.000000'
                for cell in sheet[1]:
                    cell.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
                    cell.fill = PatternFill('solid', fgColor='243746')
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                sheet.row_dimensions[1].height = 24
            detail = writer.sheets['Per_Image_Metrics']
            detail.freeze_panes = 'C2'
            detail.auto_filter.ref = detail.dimensions
            for column, width in {'A': 10, 'B': max(24, max(map(len, names)) + 2), 'C': 18, 'D': 20, 'E': 20, 'F': 18}.items():
                detail.column_dimensions[column].width = width
            for cell in detail['A'][1:]:
                cell.number_format = '0'
            writer.sheets['Summary'].column_dimensions['A'].width = 36
            writer.sheets['Summary'].column_dimensions['B'].width = 22
            writer.sheets['Summary']['B2'].number_format = '0'
        temporary.replace(path)
    finally:
        temporary.unlink(missing_ok=True)
    if not path.is_file() or path.stat().st_size == 0:
        raise RuntimeError(f'Excel export missing: {path}')
    values = dict(summary.itertuples(index=False, name=None))
    print('Excel saved successfully', flush=True)
    print(f'Path: {path.resolve()}', flush=True)
    print(f'Total rows: {len(per_image)}', flush=True)
    for label in ('Global Dice', 'Global Jaccard / Global IoU', 'Mean Dice', 'Mean Jaccard', 'Best Dice', 'Worst Dice'):
        print(f'{label}: {values[label]}', flush=True)
    return path

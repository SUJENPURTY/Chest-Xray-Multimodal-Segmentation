"""Validate QaTa annotations and prepare LanGuide CSVs without changing sources."""

from __future__ import annotations

import csv
import os
import sys
from pathlib import Path, PureWindowsPath

import numpy as np
from openpyxl import load_workbook
from PIL import Image


WORKSPACE = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(WORKSPACE))
from src.config import from_environment
from src.data.dataset import group_key, patient_group_partition, read_manifest

ANNOTATIONS = Path(os.environ.get("ANNOTATIONS_DIR", WORKSPACE / "text_annotation"))
OUTPUT = Path(os.environ["ANNOTATION_OUTPUT_DIR"]) if os.environ.get("ANNOTATION_OUTPUT_DIR") else from_environment().output_dir / "prepared_annotations"
TRAIN_VAL_FILE = ANNOTATIONS / "Train_Val_text_for_Covid19.xlsx"
TEST_FILE = ANNOTATIONS / "Test_text_for_Covid19.xlsx"
TRAIN_ID_FILE = ANNOTATIONS / "Train_ID_for_Covid19.xlsx"


def require(condition: bool, message: str) -> None:
    if not condition:
        raise ValueError(message)


def read_workbook(path: Path, required_columns: tuple[str, ...]) -> list[dict]:
    require(path.is_file(), f"Missing source workbook: {path}")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        require(len(workbook.sheetnames) == 1, f"Expected one sheet in {path.name}")
        sheet = workbook[workbook.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, None)
        require(headers is not None, f"Empty workbook: {path.name}")
        headers = list(headers)
        # These workbooks have formatted but empty columns after the data.
        while headers and headers[-1] is None:
            headers.pop()
        require(all(header is not None for header in headers),
                f"Blank header inside data columns: {path.name}")
        require(len(headers) == len(set(headers)), f"Duplicate column names: {path.name}")
        missing = set(required_columns) - set(headers)
        require(not missing, f"{path.name} lacks required columns: {sorted(missing)}")

        result = []
        for row_number, values in enumerate(rows, start=2):
            require(all(value is None for value in values[len(headers):]),
                    f"{path.name} row {row_number}: data in an unnamed column")
            row = dict(zip(headers, values[:len(headers)]))
            for column in required_columns:
                value = row[column]
                require(
                    isinstance(value, str) and bool(value.strip()),
                    f"{path.name} row {row_number}: empty or non-text {column}",
                )
            result.append(row)
        return result
    finally:
        workbook.close()


def mask_name(value: str) -> str:
    """Use only deterministic path/basename handling; never fuzzy match."""
    name = PureWindowsPath(value.strip().replace("/", "\\")).name
    require(name.startswith("mask_") and name.lower().endswith(".png"),
            f"Expected a mask_*.png filename, got {value!r}")
    return name


def image_name(mask_filename: str) -> str:
    return mask_filename.removeprefix("mask_")


def validate_unique_images(rows: list[dict], source: str) -> dict[str, dict]:
    by_name = {}
    for row in rows:
        name = mask_name(row["Image"])
        require(name not in by_name, f"Duplicate Image in {source}: {name}")
        by_name[name] = row
    return by_name


def locate_dataset() -> dict[str, tuple[Path, Path]]:
    local_default = WORKSPACE / "Project" / "LanGuideMedSeg-MICCAI2023" / "Data" / "QaTa-COV19-v2"
    default = local_default if local_default.is_dir() else WORKSPACE / "data" / "QaTa-COV19-v2"
    root = Path(os.environ.get("DATA_ROOT", default))
    paths = {
        "train": (root / "Train" / "Images", root / "Train" / "GTs"),
        "test": (root / "Test" / "Images", root / "Test" / "GTs"),
    }
    require(all(folder.is_dir() for pair in paths.values() for folder in pair),
            f"Expected Train/Images, Train/GTs, Test/Images, Test/GTs under {root}")
    return paths


def dataset_names(folders: tuple[Path, Path]) -> tuple[set[str], set[str]]:
    images, masks = folders
    return (
        {path.name for path in images.iterdir() if path.is_file() and path.suffix.lower() == ".png"},
        {path.name for path in masks.iterdir() if path.is_file() and path.suffix.lower() == ".png"},
    )


def check_pairs(rows: list[dict], folders: tuple[Path, Path], source: str) -> tuple[int, int]:
    image_names, mask_names = dataset_names(folders)
    referenced_images = set()
    missing_images = []
    missing_masks = []
    for row in rows:
        mask = mask_name(row["Image"])
        image = image_name(mask)
        referenced_images.add(image)
        if image not in image_names:
            missing_images.append(image)
        if mask not in mask_names:
            missing_masks.append(mask)
    require(not missing_images, f"{source}: {len(missing_images)} missing X-rays, e.g. {missing_images[:5]}")
    require(not missing_masks, f"{source}: {len(missing_masks)} missing masks, e.g. {missing_masks[:5]}")
    uncovered = image_names - referenced_images
    require(not uncovered, f"{source}: {len(uncovered)} dataset images lack text, e.g. {sorted(uncovered)[:5]}")
    require(len(rows) == len(image_names) == len(mask_names),
            f"{source}: annotation/image/mask counts differ")
    for row in rows:
        mask_filename = mask_name(row["Image"])
        image_path = folders[0] / image_name(mask_filename)
        mask_path = folders[1] / mask_filename
        try:
            with Image.open(image_path) as image:
                image.load()
                dimensions = image.size
            with Image.open(mask_path) as mask:
                mask.load()
                require(mask.size == dimensions, f"{source}: image/mask size mismatch: {mask_filename}")
                values = np.unique(np.asarray(mask.convert("L")))
                require(set(values.tolist()).issubset({0, 255}),
                        f"{source}: nonbinary mask: {mask_filename}, values={values[:10]}")
        except (OSError, ValueError) as exc:
            raise ValueError(f"{source}: invalid image/mask pair: {mask_filename}: {exc}") from exc
    return len(image_names), len(mask_names)


def write_csv(path: Path, columns: tuple[str, ...], rows: list[dict]) -> None:
    with path.open("w", newline="", encoding="utf-8") as file:
        writer = csv.DictWriter(file, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def reopen_csv(path: Path, columns: tuple[str, ...]) -> list[dict]:
    with path.open("r", newline="", encoding="utf-8") as file:
        reader = csv.DictReader(file)
        require(reader.fieldnames == list(columns), f"Wrong CSV columns in {path.name}: {reader.fieldnames}")
        return list(reader)


def preview(label: str, rows: list[dict], start: int) -> None:
    print(label)
    for index, row in enumerate(rows, start=start):
        print(f"  row {index} (one-based data row): {row}")


def main() -> None:
    train_val = read_workbook(TRAIN_VAL_FILE, ("Image", "Description"))
    test = read_workbook(TEST_FILE, ("Image", "Description"))
    train_id_rows = read_workbook(TRAIN_ID_FILE, ("Image",))
    train_val_by_name = validate_unique_images(train_val, TRAIN_VAL_FILE.name)
    test_by_name = validate_unique_images(test, TEST_FILE.name)
    train_ids = [mask_name(row["Image"]) for row in train_id_rows]
    train_id_set = set(train_ids)
    require(len(train_ids) == len(train_id_set), "Train_ID has duplicate Image values")
    require(not train_id_set - train_val_by_name.keys(), "Train_ID contains names outside Train_Val annotations")

    train_names = set(train_val_by_name)
    validation_names = train_names - train_id_set
    test_names = set(test_by_name)
    require(not (train_id_set & validation_names), "Train/validation leakage")
    require(not (train_id_set & test_names), "Train/test leakage")
    require(not (validation_names & test_names), "Validation/test leakage")
    require(train_id_set | validation_names == train_names, "Train/validation partition incomplete")
    require(len(validation_names) == len(train_val) - len(train_ids), "Validation complement count mismatch")

    dataset = locate_dataset()
    train_image_count, train_mask_count = check_pairs(train_val, dataset["train"], "Train_Val")
    test_image_count, test_mask_count = check_pairs(test, dataset["test"], "Test")

    # Train_ID workbook order is used for the train block. Remaining source
    # annotation order is retained within the validation block.
    train_rows = [train_val_by_name[name] for name in train_ids]
    validation_rows = [row for row in train_val if mask_name(row["Image"]) in validation_names]
    ordered_train = train_rows + validation_rows
    assert len(ordered_train) == len(train_val)

    annotation_columns = ("Image", "Description")
    to_annotation = lambda row: {key: row[key] for key in annotation_columns}
    train_csv_rows = [to_annotation(row) for row in ordered_train]
    test_csv_rows = [to_annotation(row) for row in test]
    manifest_columns = ("Image", "Split", "Description")
    grouped_train, grouped_validation = patient_group_partition(
        (row["Image"] for row in ordered_train), from_environment().seed)
    manifest_rows = [
        {"Image": row["Image"],
         "Split": "train" if group_key(row["Image"]) in grouped_train else "validation",
         "Description": row["Description"]}
        for row in ordered_train
    ] + [
        {"Image": row["Image"], "Split": "test", "Description": row["Description"]}
        for row in test
    ]

    OUTPUT.mkdir(parents=True, exist_ok=True)
    train_path = OUTPUT / "train.csv"
    test_path = OUTPUT / "test.csv"
    manifest_path = OUTPUT / "split_manifest.csv"
    write_csv(train_path, annotation_columns, train_csv_rows)
    write_csv(test_path, annotation_columns, test_csv_rows)
    write_csv(manifest_path, manifest_columns, manifest_rows)

    # Re-open every output and compare complete row content to the validated
    # source values. This also verifies that CSV quoting preserved descriptions.
    written_train = reopen_csv(train_path, annotation_columns)
    written_test = reopen_csv(test_path, annotation_columns)
    written_manifest = reopen_csv(manifest_path, manifest_columns)
    require(written_train == train_csv_rows, "train.csv differs from validated source rows")
    require(written_test == test_csv_rows, "test.csv differs from validated source rows")
    require(written_manifest == manifest_rows, "split_manifest.csv differs from validated source rows")
    boundary = len(train_ids)
    first_names = [mask_name(row["Image"]) for row in written_train[:boundary]]
    last_names = [mask_name(row["Image"]) for row in written_train[boundary:]]
    require(first_names == train_ids, "First 80% of train.csv does not follow Train_ID order")
    require(set(first_names) == train_id_set, "Train block does not equal Train_ID membership")
    require(set(last_names) == validation_names, "Validation block does not equal validation complement")
    require(len(set(first_names + last_names)) == len(written_train), "Duplicate train.csv Image")
    require(len({mask_name(row["Image"]) for row in written_test}) == len(written_test),
            "Duplicate test.csv Image")
    require(all(row["Image"].strip() and row["Description"].strip()
                for row in written_train + written_test), "Blank Image or Description in CSV")
    require(all(row["Split"] in {"train", "validation", "test"} for row in written_manifest),
            "Unknown split name in manifest")
    require(len(written_train) == len(train_val) and len(written_test) == len(test),
            "Output row count differs from source")

    dataset_root = dataset["train"][0].parent.parent
    grouped_rows = [read_manifest(train_path, dataset_root, split)
                    for split in ("train", "validation")]
    grouped_rows.append(read_manifest(test_path, dataset_root, "test"))
    grouped_keys = [{group_key(row["name"]) for row in rows} for rows in grouped_rows]
    require(not (grouped_keys[0] & grouped_keys[1] or grouped_keys[0] & grouped_keys[2]
                 or grouped_keys[1] & grouped_keys[2]), "Patient-grouped split leakage")
    require(grouped_keys[0] == grouped_train and grouped_keys[1] == grouped_validation,
            "split_manifest.csv differs from canonical patient-grouped split")
    require({row["Image"] for row in written_manifest if row["Split"] == "train"} ==
            {row["name"] for row in grouped_rows[0]}, "Manifest train rows differ from trainer")
    require({row["Image"] for row in written_manifest if row["Split"] == "validation"} ==
            {row["name"] for row in grouped_rows[1]}, "Manifest validation rows differ from trainer")

    print("train.csv rows:", len(written_train))
    print("test.csv rows:", len(written_test))
    print("Train_ID source rows:", boundary)
    print("Validation complement source rows:", len(written_train) - boundary)
    print("Train-ID matches:", len(train_id_set))
    print("Validation complement:", len(validation_names))
    print("Dataset X-rays/masks:", train_image_count, train_mask_count, test_image_count, test_mask_count)
    print("Canonical patient-grouped train/validation/test:", *(len(rows) for rows in grouped_rows))
    print("Missing images: 0; missing masks: 0; missing descriptions: 0")
    print("Duplicate filenames: 0; split leakage: 0")
    preview("First 5 train.csv rows:", written_train[:5], 1)
    preview("Last 5 train.csv rows:", written_train[-5:], len(written_train) - 4)
    preview("Train_ID/complement boundary:", written_train[max(0, boundary-3):boundary+2],
            max(1, boundary-2))
    preview("First 5 test.csv rows:", written_test[:5], 1)

    report = f"""QaTa -> LanGuide Annotation Conversion
====================================

Train_ID source rows: {len(train_ids)}
Validation complement source rows: {len(validation_names)}
Test samples: {len(test)}

train.csv rows: {len(written_train)}
test.csv rows: {len(written_test)}

Train-ID matches: {len(train_id_set)}
Validation complement: {len(validation_names)}
Missing images: 0
Missing masks: 0
Missing descriptions: 0
Duplicate filenames: 0
Split leakage: 0

train.csv order valid: YES
LanGuide annotation format: READY

Image contains the exact mask-style filename from the source workbooks.
Description text was copied without editing. The first {boundary} train.csv
data rows follow Train_ID order; the remaining {len(validation_names)} rows are
the validation complement. Output CSVs were re-opened and every row was
compared with its source value; every referenced image and mask exists.

Annotation conversion is complete. The model-ready Train/Test folder layout
was checked before these CSVs were written.
"""
    (OUTPUT / "conversion_report.txt").write_text(report, encoding="utf-8")
    print("LanGuide annotation format: READY")


if __name__ == "__main__":
    main()
